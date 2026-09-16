from io import BytesIO
from pathlib import Path
import sys
import pandas as pd
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from resume_parser import ResumeParseError, build_profile, extract_text, validate_upload
from career_intelligence import match_jobs, profile_strength, skill_gaps


def test_docx_resume_extraction_and_profile():
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("jane@example.com | Berlin | LinkedIn")
    document.add_paragraph("Education")
    document.add_paragraph("MSc Data Science")
    document.add_paragraph("Projects")
    document.add_paragraph("Built a Machine Learning project using Python, SQL, Pandas and PowerBI.")
    document.add_paragraph("Languages: German B1, English C1")
    stream = BytesIO(); document.save(stream)
    text = extract_text("resume.docx", stream.getvalue())
    profile = build_profile(text)
    assert {"Python", "SQL", "Pandas", "Power BI", "Machine Learning"} <= set(profile.skills)
    assert "German" in profile.languages and profile.education and profile.projects
    assert profile_strength(profile)["Overall profile strength"] > 0


def test_invalid_resume_uploads_are_clear():
    try: validate_upload("resume.txt", b"hello")
    except ResumeParseError as error: assert "PDF or DOCX" in str(error)
    else: raise AssertionError("Expected unsupported upload error")
    try: extract_text("resume.pdf", b"not a pdf")
    except ResumeParseError: pass
    else: raise AssertionError("Expected parsing error")


def test_matching_and_skill_gap_use_not_detected_language():
    profile = build_profile("Jane Doe\nEducation\nMSc Data Science\nProjects\nPython SQL Machine Learning project")
    jobs = pd.DataFrame([{"job_id": "1", "title": "Werkstudent Data Science", "company": "A", "location": "Berlin", "state": "BERLIN", "job_type": "PRAKTIKUM", "description": "Python SQL Tableau German B1", "salary": "", "source": "BA", "source_url": ""}])
    job_skills = pd.DataFrame([{"job_id": "1", "skill": skill} for skill in ["Python", "SQL", "Tableau"]])
    demand = pd.DataFrame([{"skill": "Python", "job_count": 5}, {"skill": "SQL", "job_count": 4}, {"skill": "Tableau", "job_count": 3}])
    matches = match_jobs(profile, jobs, job_skills)
    gaps = skill_gaps(matches, profile, demand)
    assert matches.iloc[0].match_score > 0
    assert gaps.iloc[0].resume_status == "Not detected in your resume"
    assert "Tableau" in set(gaps.skill)
