"""Local, deterministic resume extraction for PDF and DOCX uploads.

Uploaded bytes are parsed in memory only. This module does not log or persist
resume contents, and it does not call external services.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re
from typing import Any

from docx import Document
from pypdf import PdfReader

from skill_dictionary import SKILL_ALIASES, canonical_skill

MAX_RESUME_BYTES = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx"}
SECTION_HEADINGS = {"education", "experience", "work experience", "projects", "skills", "languages", "certifications"}


class ResumeParseError(ValueError):
    """An understandable upload or text-extraction failure."""


@dataclass
class ResumeProfile:
    """Structured profile extracted from evidence present in a resume."""
    name: str | None
    contact: dict[str, str]
    education: list[dict[str, str]]
    experience: list[dict[str, str]]
    projects: list[dict[str, str]]
    skills: list[str]
    skill_categories: dict[str, list[str]]
    languages: list[str]
    locations: list[str]
    target_roles: list[str]
    completeness: dict[str, bool]
    raw_text: str


def validate_upload(filename: str, content: bytes) -> None:
    """Validate supported type, size, and non-empty upload before parsing."""
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in ALLOWED_EXTENSIONS:
        raise ResumeParseError("Unsupported file type. Please upload a PDF or DOCX resume.")
    if len(content) > MAX_RESUME_BYTES:
        raise ResumeParseError("This resume is larger than 5 MB. Please upload a smaller file.")
    if not content:
        raise ResumeParseError("The uploaded file is empty. Please choose a PDF or DOCX resume.")


def extract_text(filename: str, content: bytes) -> str:
    """Extract readable text from a PDF or DOCX held in memory."""
    validate_upload(filename, content)
    try:
        if filename.lower().endswith(".pdf"):
            text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
        else:
            document = Document(BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
    except Exception as error:
        raise ResumeParseError("We couldn't read this file. Try a text-based PDF or a standard DOCX resume.") from error
    # Preserve line boundaries: they carry valuable section structure in resumes.
    text = "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines() if line.strip())
    if len(text) < 40:
        raise ResumeParseError("We couldn't extract meaningful text from this resume. Try a text-based PDF or DOCX file.")
    return text


def _find_skills(text: str) -> list[str]:
    found = []
    for skill, aliases in SKILL_ALIASES.items():
        if any(re.search(r"(?<!\w)" + re.escape(alias) + r"(?!\w)", text, re.IGNORECASE) for alias in aliases):
            found.append(skill)
    return sorted(set(found))


def _section_lines(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {"summary": []}
    active = "summary"
    for line in re.split(r"[\r\n]+", text):
        cleaned = line.strip(" :-\t")
        if cleaned.casefold() in SECTION_HEADINGS:
            active = cleaned.casefold().replace("work experience", "experience")
            sections.setdefault(active, [])
        elif cleaned:
            sections.setdefault(active, []).append(cleaned)
    return {name: " ".join(lines) for name, lines in sections.items()}


def _categorize(skills: list[str]) -> dict[str, list[str]]:
    groups = {
        "Programming": {"Python", "R", "Java", "C#", ".NET"},
        "Data analytics": {"SQL", "Pandas", "NumPy", "Excel", "Power BI", "Tableau", "Statistics"},
        "Data science and ML": {"Machine Learning", "Deep Learning", "scikit-learn", "Data Science", "NLP"},
        "AI": {"Artificial Intelligence", "Generative AI", "Large Language Models", "RAG", "Embeddings"},
        "Cloud and DevOps": {"AWS", "Azure", "Docker", "Git", "GitLab", "Cloud"},
        "Data engineering": {"Spark", "PySpark", "Hive", "Data Engineering"},
    }
    return {group: sorted(set(skills) & values) for group, values in groups.items() if set(skills) & values}


def build_profile(text: str) -> ResumeProfile:
    """Build a transparent profile using only patterns evidenced in the resume text."""
    sections = _section_lines(text)
    lines = [line.strip() for line in re.split(r"[\r\n]+", text) if line.strip()]
    first_line = lines[0] if lines and len(lines[0]) < 80 and "@" not in lines[0] else None
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone = re.search(r"(?:\+?\d[\d\s()./-]{7,}\d)", text)
    linkedin = re.search(r"(?:https?://)?(?:www\.)?linkedin\.com/[^\s|]+", text, re.I)
    github = re.search(r"(?:https?://)?(?:www\.)?github\.com/[^\s|]+", text, re.I)
    location = re.search(r"\b(?:Berlin|Potsdam|Munich|München|Hamburg|Frankfurt|Cologne|Köln|Dresden|Bremen)\b", text, re.I)
    languages = []
    for language in ("German", "Deutsch", "English", "Englisch", "French", "Spanish"):
        if re.search(r"(?<!\w)" + language + r"(?!\w)", text, re.I):
            languages.append("German" if language == "Deutsch" else "English" if language == "Englisch" else language)
    education_text = sections.get("education", "")
    degree = re.search(r"\b(?:MSc|M\.Sc\.|Master(?:'s)?|BSc|B\.Sc\.|Bachelor)\b[^.]{0,120}", education_text or text, re.I)
    education = [{"evidence": degree.group(0).strip()}] if degree else []
    experience = [{"evidence": sections["experience"]}] if sections.get("experience") else []
    projects = [{"evidence": sections["projects"]}] if sections.get("projects") else []
    target_roles = [role for role in ("Data Scientist", "Data Analyst", "Data Engineer", "Machine Learning Engineer", "Software Engineer") if role.casefold() in text.casefold()]
    skills = _find_skills(text)
    contact = {key: value for key, value in {"email": email.group(0) if email else "", "phone": phone.group(0) if phone else "", "linkedin": linkedin.group(0) if linkedin else "", "github": github.group(0) if github else ""}.items() if value}
    completeness = {"contact": bool(contact), "education": bool(education), "experience": bool(experience), "projects": bool(projects), "skills": bool(skills), "languages": bool(languages), "location": bool(location)}
    return ResumeProfile(first_line, contact, education, experience, projects, skills, _categorize(skills), sorted(set(languages)), [location.group(0) if location else ""] if location else [], target_roles, completeness, text)
